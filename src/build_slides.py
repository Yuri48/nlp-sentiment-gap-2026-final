"""
build_slides.py — 6-8 minute summary presentation (python-pptx), built from presentation/slides.md.

slides.md dialect: a slide starts with '## Title'; lines '- ' are bullets; a line '![](figures/x.png)'
places the image on the right half (or full width if the slide has no bullets); a line starting with
'> ' becomes the speaker note; '%% key: value' sets metadata (title slide fields).
Usage: python src/build_slides.py presentation/slides.md presentation/Presentation.pptx
"""
import os, re, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
src, dst = sys.argv[1], sys.argv[2]
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
NAVY = RGBColor(0x1f, 0x3a, 0x5f); GREY = RGBColor(0x55, 0x55, 0x55); RED = RGBColor(0xc0, 0x39, 0x2b)

def add_text(slide, left, top, width, height, text, size=18, bold=False, color=None, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(left, top, width, height); tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; r = p.add_run(); r.text = text; r.font.size = Pt(size); r.font.bold = bold; r.font.name = 'Calibri'
    if color: r.font.color.rgb = color
    p.alignment = align
    return tb

def add_bullets(slide, left, top, width, height, bullets, size=16):
    tb = slide.shapes.add_textbox(left, top, width, height); tf = tb.text_frame; tf.word_wrap = True
    first = True
    for b in bullets:
        lvl = 1 if b.startswith('  ') else 0
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        txt = b.strip()
        for m in re.finditer(r'(\*\*.+?\*\*|[^*]+)', txt):
            r = p.add_run(); t = m.group(0)
            if t.startswith('**'):
                r.text = t[2:-2]; r.font.bold = True
            else:
                r.text = t
            r.font.size = Pt(size - 2 * lvl); r.font.name = 'Calibri'; r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
        p.level = lvl; p.space_after = Pt(6)
        pPr = p._p.get_or_add_pPr()
        from pptx.oxml.ns import qn
        from pptx.oxml import parse_xml
        bu = parse_xml(f'<a:buChar xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" char="{"•" if lvl == 0 else "–"}"/>')
        pPr.insert(0, parse_xml('<a:buFont xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="Arial"/>')); pPr.append(bu)
        pPr.set('marL', str(int(Emu(Inches(0.3 + 0.3 * lvl))))); pPr.set('indent', str(-int(Emu(Inches(0.25)))))

meta, slides = {}, []
cur = None
for ln in open(src, encoding='utf-8').read().splitlines():
    if ln.startswith('%%'):
        k, v = ln[2:].split(':', 1); meta[k.strip()] = v.strip()
    elif ln.startswith('## '):
        cur = {'title': ln[3:].strip(), 'bullets': [], 'images': [], 'notes': []}; slides.append(cur)
    elif cur is None:
        continue
    elif ln.startswith('!['):
        cur['images'].append(re.search(r'\((.*?)\)', ln).group(1))
    elif ln.startswith('> '):
        cur['notes'].append(ln[2:])
    elif re.match(r'^\s*- ', ln):
        cur['bullets'].append(re.sub(r'^(\s*)- ', lambda m: '  ' if m.group(1) else '', ln))

# title slide
s = prs.slides.add_slide(BLANK)
s.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(7.5)).fill.solid(); s.shapes[-1].fill.fore_color.rgb = NAVY; s.shapes[-1].line.fill.background()
add_text(s, Inches(0.8), Inches(2.2), Inches(11.5), Inches(1.6), meta.get('title', 'Title'), size=34, bold=True, color=RGBColor(0xff, 0xff, 0xff))
add_text(s, Inches(0.8), Inches(3.9), Inches(11.5), Inches(1.0), meta.get('subtitle', ''), size=20, color=RGBColor(0xdd, 0xe6, 0xf0))
add_text(s, Inches(0.8), Inches(5.6), Inches(11.5), Inches(1.0), meta.get('author', ''), size=16, color=RGBColor(0xdd, 0xe6, 0xf0))
for sd in slides:
    s = prs.slides.add_slide(BLANK)
    s.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.12)).fill.solid(); s.shapes[-1].fill.fore_color.rgb = NAVY; s.shapes[-1].line.fill.background()
    add_text(s, Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.9), sd['title'], size=26, bold=True, color=NAVY)
    has_img, has_b = bool(sd['images']), bool(sd['bullets'])
    if has_b:
        add_bullets(s, Inches(0.5), Inches(1.3), Inches(6.1) if has_img else Inches(12.3), Inches(5.6), sd['bullets'], size=16 if has_img else 18)
    if has_img:
        n = len(sd['images'])
        for k, img in enumerate(sd['images']):
            path = os.path.join(ROOT, img)
            if not os.path.exists(path):
                continue
            if has_b:
                left, top, width = Inches(6.8), Inches(1.3 + k * (5.6 / n)), Inches(6.1)
                s.shapes.add_picture(path, left, top, width=width) if n == 1 else s.shapes.add_picture(path, left, top, height=Inches(5.4 / n))
            else:
                s.shapes.add_picture(path, Inches(0.7 + k * (12 / n)), Inches(1.4), width=Inches(11.8 / n))
    if sd['notes']:
        s.notes_slide.notes_text_frame.text = ' '.join(sd['notes'])
    add_text(s, Inches(0.5), Inches(7.05), Inches(12), Inches(0.4), meta.get('footer', ''), size=10, color=GREY)
prs.save(dst); print('saved', dst, len(prs.slides), 'slides')
