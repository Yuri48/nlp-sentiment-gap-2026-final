"""
md_to_docx.py — Build the Word report from report/report.md (a restricted Markdown dialect).

Supported: # / ## / ### headings, paragraphs, '- ' bullets, '1. ' numbered items, pipe tables,
images ![caption](relative/path.png){width=6.0}, **bold**, *italic*, `code`, horizontal rule '---'
(page break), and a line starting with '> ' rendered as an indented note.
Usage: python src/md_to_docx.py report/report.md report/Report.docx
"""
import os, re, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
src, dst = sys.argv[1], sys.argv[2]
doc = Document()
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
st.paragraph_format.space_after = Pt(5); st.paragraph_format.line_spacing = 1.1
for s in doc.sections:
    s.left_margin = s.right_margin = Inches(0.95); s.top_margin = s.bottom_margin = Inches(0.85)
for lvl, size in [(1, 16), (2, 13), (3, 11.5)]:
    h = doc.styles[f'Heading {lvl}']; h.font.name = 'Calibri'; h.font.size = Pt(size); h.font.color.rgb = RGBColor(0x1f, 0x3a, 0x5f)

def add_runs(par, text):
    pos = 0
    for m in re.finditer(r'(\*\*.+?\*\*|\*.+?\*|`.+?`)', text):
        if m.start() > pos:
            par.add_run(text[pos:m.start()])
        tok = m.group(0)
        if tok.startswith('**'):
            par.add_run(tok[2:-2]).bold = True
        elif tok.startswith('`'):
            r = par.add_run(tok[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(9.5)
        else:
            par.add_run(tok[1:-1]).italic = True
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:])

def set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexcolor); tcPr.append(shd)

def add_table(lines):
    rows = [[c.strip() for c in l.strip().strip('|').split('|')] for l in lines if not re.match(r'^\s*\|?\s*:?-{2,}', l)]
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol); t.style = 'Light Grid Accent 1'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, r in enumerate(rows):
        for j in range(ncol):
            cell = t.cell(i, j); cell.text = ''
            p = cell.paragraphs[0]; add_runs(p, r[j] if j < len(r) else '')
            for run in p.runs:
                run.font.size = Pt(8.5); run.bold = run.bold or (i == 0)
            if i == 0:
                set_cell_bg(cell, 'DCE6F1')
    doc.add_paragraph()

lines = open(src, encoding='utf-8').read().splitlines()
i = 0
while i < len(lines):
    ln = lines[i]
    if ln.startswith('|'):
        j = i
        while j < len(lines) and lines[j].startswith('|'):
            j += 1
        add_table(lines[i:j]); i = j; continue
    m = re.match(r'^!\[(.*?)\]\((.*?)\)(\{width=([\d.]+)\})?', ln)
    if m:
        path = os.path.join(ROOT, m.group(2)) if not os.path.isabs(m.group(2)) else m.group(2)
        if os.path.exists(path):
            doc.add_picture(path, width=Inches(float(m.group(4)) if m.group(4) else 6.3))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cap.add_run(m.group(1)); r.italic = True; r.font.size = Pt(9)
        i += 1; continue
    if ln.startswith('# '):
        doc.add_heading(ln[2:].strip(), level=1)
    elif ln.startswith('## '):
        doc.add_heading(ln[3:].strip(), level=2)
    elif ln.startswith('### '):
        doc.add_heading(ln[4:].strip(), level=3)
    elif ln.strip() == '---':
        doc.add_page_break()
    elif re.match(r'^\s*- ', ln):
        p = doc.add_paragraph(style='List Bullet'); add_runs(p, re.sub(r'^\s*- ', '', ln))
    elif re.match(r'^\s*\d+\. ', ln):
        p = doc.add_paragraph(style='List Number'); add_runs(p, re.sub(r'^\s*\d+\. ', '', ln))
    elif ln.startswith('> '):
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.4); add_runs(p, ln[2:])
        for r in p.runs:
            r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
    elif ln.strip() == '':
        pass
    else:
        # merge consecutive non-empty plain lines into one paragraph
        j = i; buf = []
        while j < len(lines) and lines[j].strip() != '' and not re.match(r'^(#|\||!\[|- |\d+\. |> |---)', lines[j]):
            buf.append(lines[j].strip()); j += 1
        p = doc.add_paragraph(); add_runs(p, ' '.join(buf)); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        i = j; continue
    i += 1
doc.save(dst)
print('saved', dst)
